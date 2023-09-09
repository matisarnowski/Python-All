#! /bin/python3

import openpyxl
import os
from openpyxl.chart import BarChart, Series, Reference
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
sheet = openpyxl.load_workbook(os.path.join(os.getcwd(), 'example.xlsx'))
Sheet1 = sheet["Sheet1"]

list_with_cells = []

for row in Sheet1:
    temp_list = []
    for cell in row:
        temp_list.append(cell.value)
    list_with_cells.append(temp_list)

print(list_with_cells)

for row in list_with_cells:
    ws.append(row)

chart = BarChart(type="col")
chart.type = "col"
chart.title = "Ilość owoców na targu."
chart.y_axis.title = 'Sztuk'
chart.x_axis.title = 'Owoc'
chart.legend = None

data = Reference(ws, min_col=3, min_row=2, max_row=8, max_col=3)
categories = Reference(ws, min_col=2, min_row=2, max_row=8, max_col=2)

chart.add_data(data)
chart.set_categories(categories)

ws.add_chart(chart, "E1")
wb.save("example1.xlsx")
